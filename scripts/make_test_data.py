"""Generate a fictional company's files for testing Sirgal.

Some files contain sensitive data (salaries, SSNs, card numbers, passwords),
some are harmless. manifest.json records what each file contains, how it is
meant to be shared, and which detection layer should catch it:

    "rules"       structured data that patterns and checksums can find
    "model"       sensitive details in written text, which needs a local model
    "unreadable"  a file with no text to read (like a scanned image), which
                  Sirgal must report as UNKNOWN, never OK

Some harmless files are full of people's names on purpose. Almost every
document mentions people, so names alone must not make a file risky.

Everything here is fictional, made up by Faker. The files come in the formats
companies actually use: text, CSV, PDF, Word and Excel.

Usage:
    python scripts/make_test_data.py
    python scripts/make_test_data.py --employees 50 --seed 7
"""

import argparse
import csv
import json
import textwrap
from pathlib import Path

from docx import Document
from faker import Faker
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

COMPANY = "Acme Robotics"
DOMAIN = "acme.example"  # .example is reserved, so it can never be a real domain

# How each file is meant to be shared once uploaded to Google Drive.
PRIVATE = "private"          # only the owner
COMPANY_WIDE = "company"     # everyone in the organization
PUBLIC_LINK = "public_link"  # anyone with the link


def write_csv(path, header, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(header)
        writer.writerows(rows)


def write_text(path, text):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def write_docx(path, paragraphs):
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    doc.save(path)


def write_xlsx(path, sheets):
    """sheets: {sheet name: list of rows}, in order."""
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    for i, (name, rows) in enumerate(sheets.items()):
        sheet = workbook.active if i == 0 else workbook.create_sheet()
        sheet.title = name
        for row in rows:
            sheet.append(row)
    workbook.save(path)


def write_pdf(path, lines):
    """A simple one-page PDF with real, selectable text."""
    path.parent.mkdir(parents=True, exist_ok=True)
    page = canvas.Canvas(str(path), pagesize=A4, invariant=1)
    y = 800
    for line in lines:
        for part in textwrap.wrap(line, 90) or [""]:
            page.drawString(60, y, part)
            y -= 18
    page.showPage()
    page.save()


def write_scanned_pdf(path):
    """A PDF with shapes only and no text layer, like a scanned ID card."""
    path.parent.mkdir(parents=True, exist_ok=True)
    page = canvas.Canvas(str(path), pagesize=A4, invariant=1)
    page.setFillGray(0.85)
    page.roundRect(60, 560, 470, 240, 12, fill=1, stroke=0)   # the card
    page.setFillGray(0.6)
    page.rect(80, 600, 120, 160, fill=1, stroke=0)            # the photo
    for i in range(6):                                        # printed lines
        page.rect(220, 740 - i * 24, 280 - i * 20, 8, fill=1, stroke=0)
    page.showPage()
    page.save()


def make_employees(fake, count):
    people = []
    for _ in range(count):
        first, last = fake.first_name(), fake.last_name()
        people.append({
            "name": f"{first} {last}",
            "email": f"{first}.{last}@{DOMAIN}".lower(),
            "job": fake.job(),
            "ssn": fake.ssn(),
            "dob": fake.date_of_birth(minimum_age=22, maximum_age=64).isoformat(),
            "salary": fake.random_int(45_000, 180_000, step=500),
            "iban": fake.iban(),
        })
    return people


def main():
    parser = argparse.ArgumentParser(description="Generate fictional company files for testing Sirgal.")
    parser.add_argument("--out", default="test-data", help="output folder (default: test-data)")
    parser.add_argument("--employees", type=int, default=25, help="number of employees (default: 25)")
    parser.add_argument("--customers", type=int, default=40, help="number of customers (default: 40)")
    parser.add_argument("--seed", type=int, default=42, help="same seed gives the same files every time")
    args = parser.parse_args()

    Faker.seed(args.seed)
    fake = Faker("en_US")
    root = Path(args.out) / "acme-robotics"
    manifest = []

    def record(rel_path, sharing, contains, sensitive=None, layer=None):
        if sensitive is None:
            sensitive = bool(contains)
        manifest.append({
            "path": rel_path,
            "sharing": sharing,
            "sensitive": sensitive,
            "contains": contains,
            "layer": layer if layer else ("rules" if sensitive else None),
        })

    staff = make_employees(fake, args.employees)

    # 1. Salary sheet shared with anyone who has the link. The worst case.
    write_csv(
        root / "HR/salaries_2026.csv",
        ["name", "email", "job_title", "annual_salary_usd", "bank_iban"],
        [[p["name"], p["email"], p["job"], p["salary"], p["iban"]] for p in staff],
    )
    record("HR/salaries_2026.csv", PUBLIC_LINK, ["person_name", "email", "salary", "iban"])

    # 2. Employee records visible to the whole company. Classic Copilot risk.
    write_csv(
        root / "HR/employee_records.csv",
        ["name", "email", "ssn", "date_of_birth"],
        [[p["name"], p["email"], p["ssn"], p["dob"]] for p in staff],
    )
    record("HR/employee_records.csv", COMPANY_WIDE, ["person_name", "email", "ssn", "date_of_birth"])

    # 3. Customer list with card numbers, correctly kept private.
    write_csv(
        root / "Sales/customers.csv",
        ["company", "contact", "email", "phone", "card_number"],
        [
            [fake.company(), fake.name(), fake.email(), fake.phone_number(), fake.credit_card_number()]
            for _ in range(args.customers)
        ],
    )
    record("Sales/customers.csv", PRIVATE, ["person_name", "email", "phone", "credit_card"])

    # 4. Shared passwords file, visible to the whole company.
    services = ["AWS root", "Stripe", "Office Wi-Fi", "QuickBooks", "GitHub admin", "Payroll portal"]
    lines = [f"{s}: {fake.user_name()} / {fake.password(length=14)}" for s in services]
    write_text(root / "IT/passwords.txt", "Shared logins, do not forward\n\n" + "\n".join(lines) + "\n")
    record("IT/passwords.txt", COMPANY_WIDE, ["password"])

    # 5. Harmless file shared publicly. Sirgal should NOT flag this.
    ideas = "\n".join(f"- {fake.catch_phrase()}" for _ in range(10))
    write_text(root / "Marketing/blog_ideas.txt", f"{COMPANY} blog ideas\n\n{ideas}\n")
    record("Marketing/blog_ideas.txt", PUBLIC_LINK, [])

    # 6. Harmless file shared company-wide. Also should NOT be flagged.
    write_text(
        root / "General/team_lunch.txt",
        f"Team lunch on Friday at 1pm.\nPlace: {fake.city()} Grill\n{fake.paragraph(nb_sentences=3)}\n",
    )
    record("General/team_lunch.txt", COMPANY_WIDE, [])

    # ---- Written documents. Rules can't catch these; a local model should. ----

    # 7. Exit interview with a home address, visible to the whole company.
    leaver = fake.random_element(staff)
    leaver_first = leaver["name"].split()[0]
    reason = fake.random_element([
        "a long commute and no option to work remotely",
        "wanting to move into a management role",
        "moving closer to family",
    ])
    address = f"{fake.street_address()}, {fake.city()}, {fake.state_abbr()} {fake.zipcode()}"
    write_text(
        root / "HR/exit_interview_notes.txt",
        f"Exit interview notes\n\n"
        f"Employee: {leaver['name']}, {leaver['job']}\n\n"
        f"{leaver_first} said the main reason for leaving was {reason}. "
        f"{leaver_first} was positive about the team but felt the role had stopped growing.\n\n"
        f"Send the final documents to {leaver_first}'s home address: {address}.\n",
    )
    record("HR/exit_interview_notes.txt", COMPANY_WIDE, ["person_name", "home_address"], layer="model")

    # 8. Medical leave notes, visible to the whole company.
    conditions = [
        "recovery after knee surgery",
        "treatment for anxiety",
        "complications during pregnancy",
        "ongoing chemotherapy",
        "a back injury",
    ]
    leave_lines = []
    for person in fake.random_elements(staff, length=4, unique=True):
        first = person["name"].split()[0]
        leave_lines.append(
            f"{person['name']} is on medical leave for {fake.random_element(conditions)}. "
            f"{first} expects to return in {fake.random_int(2, 10)} weeks."
        )
    write_text(root / "HR/medical_leave_notes.txt", "Leave tracker\n\n" + "\n".join(leave_lines) + "\n")
    record("HR/medical_leave_notes.txt", COMPANY_WIDE, ["person_name", "health_info"], layer="model")

    # 9. Meeting notes full of names but nothing sensitive. Must stay OK.
    attendees = [p["name"] for p in fake.random_elements(staff, length=4, unique=True)]
    firsts = [n.split()[0] for n in attendees]
    write_text(
        root / "General/weekly_sync_notes.txt",
        f"Weekly product sync\n\nAttendees: {', '.join(attendees)}\n\n"
        f"{firsts[0]} walked through the roadmap for next quarter. "
        f"{firsts[1]} will follow up with the design team about the new onboarding screens. "
        f"{firsts[2]} and {firsts[3]} are pairing on the reporting bug.\n\n"
        f"Next sync is on Thursday.\n",
    )
    record("General/weekly_sync_notes.txt", COMPANY_WIDE, ["person_name"], sensitive=False)

    # 10. Public customer quote with a name and a city. Public on purpose, must stay OK.
    write_text(
        root / "Marketing/customer_quote.txt",
        f"Customer quote for the website\n\n"
        f"\"{fake.catch_phrase()} is exactly what our warehouse needed.\" "
        f"said {fake.name()}, operations lead at {fake.company()} in {fake.city()}.\n",
    )
    record("Marketing/customer_quote.txt", PUBLIC_LINK, ["person_name", "city"], sensitive=False)

    # 11. Office move note with a real business address. Must stay OK:
    # company addresses are public and show up in many ordinary files.
    office = f"{fake.building_number()} West Madison Street, Chicago, IL 60602"
    write_text(
        root / "General/office_move.txt",
        f"Office move\n\nFrom the first of next month we're moving to {office}. "
        f"Parking is on the north side of the building. "
        f"Questions go to the facilities channel.\n",
    )
    record("General/office_move.txt", COMPANY_WIDE, ["business_address", "city"], sensitive=False)

    # 12. General leave policy that mentions medical leave but no people. Must stay OK.
    write_text(
        root / "HR/leave_policy.txt",
        "Leave policy\n\n"
        "Employees can take up to 12 weeks of medical leave per year. "
        "Tell your manager as early as you can, and HR will help with the paperwork. "
        "Parental leave and holidays are covered in the handbook.\n",
    )
    record("HR/leave_policy.txt", COMPANY_WIDE, ["health_topic"], sensitive=False)

    # ---- PDF, Word and Excel files. ----

    # 13. Payroll workbook. The first sheet is harmless; the second has salaries
    # and bank details, so a reader that only checks the first sheet misses it.
    departments = ["Engineering", "Sales", "Operations", "Support"]
    write_xlsx(root / "HR/payroll_q3.xlsx", {
        "Summary": [["Department", "Headcount", "Total cost USD"]]
                   + [[d, fake.random_int(3, 12), fake.random_int(300, 1500) * 1000] for d in departments],
        "Detail": [["Name", "Annual salary USD", "Bank IBAN"]]
                  + [[p["name"], p["salary"], p["iban"]] for p in staff],
    })
    record("HR/payroll_q3.xlsx", COMPANY_WIDE, ["person_name", "salary", "iban"])

    # 14. Offer letter with a salary and a home address.
    hire = fake.random_element(staff)
    hire_first = hire["name"].split()[0]
    write_docx(root / "HR/offer_letter.docx", [
        "Offer of employment",
        hire["name"],
        f"{fake.street_address()}, {fake.city()}, {fake.state_abbr()} {fake.zipcode()}",
        f"Dear {hire_first},",
        f"We are pleased to offer you the position of {hire['job']} at {COMPANY}. "
        f"Your annual salary will be ${hire['salary']:,}, paid monthly.",
        "Please sign and return this letter within five working days.",
    ])
    record("HR/offer_letter.docx", COMPANY_WIDE, ["person_name", "home_address", "salary"])

    # 15. Sick note as a PDF. Only a model can tell this is health information.
    patient = fake.random_element(staff)
    write_pdf(root / "HR/sick_note.pdf", [
        "Medical certificate",
        "",
        f"This is to confirm that {patient['name']} was seen at our clinic this week.",
        f"Diagnosis: {fake.random_element(conditions)}.",
        f"Recommended rest: {fake.random_int(3, 14)} days.",
    ])
    record("HR/sick_note.pdf", COMPANY_WIDE, ["person_name", "health_info"], layer="model")

    # 16. Supplier invoice, public. Business addresses and totals only. Must stay OK.
    write_pdf(root / "Finance/vendor_invoice.pdf", [
        f"Invoice INV-{fake.random_int(10000, 99999)}",
        "",
        f"From: {fake.company()}, {fake.street_address()}, {fake.city()}",
        f"Bill to: {COMPANY}, {office}",
        "",
        "Servo motors x 40          4,800.00",
        "Control boards x 25        7,600.00",
        "Total due                 12,400.00 USD",
        "",
        "Payment due within 30 days.",
    ])
    record("Finance/vendor_invoice.pdf", PUBLIC_LINK, ["business_address"], sensitive=False)

    # 17. Holiday calendar. Must stay OK.
    write_xlsx(root / "General/holiday_calendar.xlsx", {
        "2026": [["Date", "Holiday"], ["2026-01-01", "New Year's Day"], ["2026-07-04", "Independence Day"],
                 ["2026-11-26", "Thanksgiving"], ["2026-12-25", "Christmas Day"]],
    })
    record("General/holiday_calendar.xlsx", COMPANY_WIDE, [], sensitive=False)

    # 18. Scanned passport: a PDF with no text at all. Sirgal can't read it, so a
    # shared copy must show as UNKNOWN, never OK.
    write_scanned_pdf(root / "HR/passport_scan.pdf")
    record("HR/passport_scan.pdf", COMPANY_WIDE, ["passport"], layer="unreadable")

    manifest_data = {"company": COMPANY, "seed": args.seed, "files": manifest}
    write_text(root / "manifest.json", json.dumps(manifest_data, indent=2) + "\n")

    risky = sum(1 for m in manifest if m["sensitive"] and m["sharing"] != PRIVATE)
    print(f"Created {len(manifest)} files in {root}/")
    model_only = sum(1 for m in manifest if m["layer"] == "model")
    unreadable = sum(1 for m in manifest if m["layer"] == "unreadable")
    print(f"{risky} are sensitive and overshared ({model_only} need the model layer, "
          f"{unreadable} can't be read). See manifest.json for details.")


if __name__ == "__main__":
    main()