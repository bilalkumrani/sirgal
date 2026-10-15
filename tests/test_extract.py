"""Tests for reading PDF, Word, Excel and PowerPoint files in memory."""

import io
import zipfile

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas

from sirgal.extract import DOCX, MAX_CHARS, MAX_UNZIPPED_BYTES, PDF, PPTX, XLSX, extract_text, kind_for


def make_pdf(lines=(), shapes_only=False, password=None):
    buffer = io.BytesIO()
    page = canvas.Canvas(buffer, invariant=1)
    if password:
        from reportlab.lib import pdfencrypt

        page = canvas.Canvas(buffer, invariant=1, encrypt=pdfencrypt.StandardEncryption(password, canPrint=0))
    if shapes_only:
        page.rect(80, 600, 120, 160, fill=1)
    for i, line in enumerate(lines):
        page.drawString(60, 800 - i * 18, line)
    page.showPage()
    page.save()
    return buffer.getvalue()


def make_docx(paragraphs=(), table=None):
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def make_xlsx(sheets):
    workbook = Workbook()
    for i, (name, rows) in enumerate(sheets.items()):
        sheet = workbook.active if i == 0 else workbook.create_sheet()
        sheet.title = name
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_pptx(slides):
    """slides: a list of dicts with optional "text" (lines), "table" (rows) and "notes"."""
    deck = Presentation()
    for spec in slides:
        slide = deck.slides.add_slide(deck.slide_layouts[6])  # blank layout
        if spec.get("text"):
            box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text_frame
            box.text = spec["text"][0]
            for line in spec["text"][1:]:
                box.add_paragraph().text = line
        if spec.get("table"):
            rows = spec["table"]
            table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(1), Inches(3), Inches(8), Inches(2)).table
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = spec["notes"]
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def test_kind_from_mime_or_name():
    assert kind_for(mime="application/pdf") == PDF
    assert kind_for(name="Offer Letter.DOCX") == DOCX
    assert kind_for(name="payroll.xlsx") == XLSX
    assert kind_for(name="Board Update.PPTX") == PPTX
    assert kind_for(mime="application/vnd.openxmlformats-officedocument.presentationml.presentation") == PPTX
    assert kind_for(mime="image/png", name="photo.png") is None


def test_pdf_text_is_read():
    text = extract_text(make_pdf(["SSN 251-29-2287"]), PDF)
    assert "251-29-2287" in text


def test_scanned_pdf_has_no_text():
    assert extract_text(make_pdf(shapes_only=True), PDF) is None


def test_password_protected_pdf_is_not_readable():
    assert extract_text(make_pdf(["secret"], password="hunter2"), PDF) is None


def test_damaged_files_never_crash_the_scan():
    for kind in (PDF, DOCX, XLSX, PPTX):
        assert extract_text(b"this is not a real file", kind) is None


def test_word_paragraphs_and_tables_are_read():
    data = make_docx(["Dear Ali,"], table=[["Salary", "95,000"], ["IBAN", "GB82WEST12345698765432"]])
    text = extract_text(data, DOCX)
    assert "Dear Ali," in text
    assert "95,000" in text and "GB82WEST12345698765432" in text


def test_every_excel_sheet_is_read_not_just_the_first():
    data = make_xlsx({"Summary": [["Team", "Headcount"], ["Sales", 4]],
                      "Detail": [["Name", "IBAN"], ["Ali", "GB82WEST12345698765432"]]})
    text = extract_text(data, XLSX)
    assert "Sales" in text
    assert "GB82WEST12345698765432" in text


def test_empty_workbook_counts_as_no_text():
    assert extract_text(make_xlsx({"Sheet1": []}), XLSX) is None


def test_powerpoint_slides_tables_and_notes_are_read():
    data = make_pptx([
        {"text": ["Compensation review", "Q3 2026"]},
        {"table": [["Name", "Salary"], ["Ali", "95,000"]], "notes": "IBAN GB82WEST12345698765432"},
    ])
    text = extract_text(data, PPTX)
    assert "Compensation review\nQ3 2026" in text
    assert "Ali,95,000" in text  # a table row reads as one line, like Word tables
    assert "GB82WEST12345698765432" in text  # speaker notes are checked too


def test_slides_are_read_in_order_past_slide_nine():
    data = make_pptx([{"text": [f"slide {i}"]} for i in range(1, 12)])
    text = extract_text(data, PPTX)
    assert text.index("slide 2\n") < text.index("slide 10\n")


def test_deck_without_text_counts_as_no_text():
    assert extract_text(make_pptx([{}, {}]), PPTX) is None


def test_powerpoint_text_stops_at_the_limit():
    line = "x" * 100_000
    data = make_pptx([{"text": [line] * 10} for _ in range(5)])
    assert len(extract_text(data, PPTX)) == MAX_CHARS


def test_powerpoint_zip_bomb_is_not_unpacked():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("ppt/slides/slide1.xml", b"<a/>" + b" " * MAX_UNZIPPED_BYTES)
    assert len(buffer.getvalue()) < 1_000_000  # small to download, huge once unpacked
    assert extract_text(buffer.getvalue(), PPTX) is None


def test_powerpoint_entity_expansion_is_refused():
    slide = (b'<?xml version="1.0"?><!DOCTYPE p [<!ENTITY a "aaaaaaaaaa">]>'
             b'<p xmlns="http://schemas.openxmlformats.org/drawingml/2006/main"><t>&a;</t></p>')
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("ppt/slides/slide1.xml", slide)
    assert extract_text(buffer.getvalue(), PPTX) is None
