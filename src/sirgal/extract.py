"""Pull plain text out of PDF, Word, Excel and PowerPoint files, entirely in memory.

Files arrive as bytes and never touch the disk. When a file has no text to
read (a scanned image, a password-protected PDF, a damaged file), this returns
None, so the scan reports it as UNKNOWN instead of calling it safe.
"""

import io
import re
import warnings
import zipfile
from xml.etree import ElementTree

PDF = "pdf"
DOCX = "docx"
XLSX = "xlsx"
PPTX = "pptx"

MIME_TYPES = {
    "application/pdf": PDF,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": DOCX,
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": XLSX,
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": PPTX,
}
SUFFIXES = {".pdf": PDF, ".docx": DOCX, ".xlsx": XLSX, ".pptx": PPTX}

MAX_CHARS = 2_000_000  # stop collecting text past this, to keep memory bounded
MAX_UNZIPPED_BYTES = 100 * 1024 * 1024  # a zip bomb unpacks far past this; treat it as unreadable


def kind_for(mime=None, name=None):
    """The document kind for a MIME type or file name, or None if not supported."""
    if mime in MIME_TYPES:
        return MIME_TYPES[mime]
    if name:
        for suffix, kind in SUFFIXES.items():
            if name.lower().endswith(suffix):
                return kind
    return None


def extract_text(data: bytes, kind: str):
    """Return the text in a PDF, Word, Excel or PowerPoint file, or None if there is none to read."""
    readers = {PDF: _pdf_text, DOCX: _docx_text, XLSX: _xlsx_text, PPTX: _pptx_text}
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")  # damaged files trigger library warnings
            text = readers[kind](io.BytesIO(data))
    except Exception:  # a damaged or unusual file must never stop the scan
        return None
    text = text[:MAX_CHARS]
    return text if text.strip() else None


def _pdf_text(stream):
    from pypdf import PdfReader

    reader = PdfReader(stream)
    if reader.is_encrypted and not reader.decrypt(""):
        return ""  # password-protected: nothing we can read
    parts, size = [], 0
    for page in reader.pages:
        part = page.extract_text() or ""
        parts.append(part)
        size += len(part)
        if size > MAX_CHARS:
            break
    return "\n".join(parts)


def _docx_text(stream):
    from docx import Document

    doc = Document(stream)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:  # offer letters and forms often keep details in tables
        for row in table.rows:
            parts.append(",".join(cell.text for cell in row.cells))
    for section in doc.sections:
        for block in (section.header, section.footer):
            parts.extend(p.text for p in block.paragraphs)
    return "\n".join(parts)


def _xlsx_text(stream):
    from openpyxl import load_workbook

    workbook = load_workbook(stream, read_only=True, data_only=True)
    parts, size = [], 0
    try:
        for sheet in workbook.worksheets:  # every sheet, not just the first
            for row in sheet.iter_rows(values_only=True):
                line = ",".join("" if v is None else str(v) for v in row)
                parts.append(line)
                size += len(line)
                if size > MAX_CHARS:
                    return "\n".join(parts)
    finally:
        workbook.close()
    return "\n".join(parts)


_DRAWING = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_SLIDE_PART = re.compile(r"ppt/(slides/slide|notesSlides/notesSlide)(\d+)\.xml")


def _pptx_text(stream):
    """Text from every slide, its tables and its speaker notes, read with the standard library."""
    with zipfile.ZipFile(stream) as archive:
        parts = []
        for info in archive.infolist():
            match = _SLIDE_PART.fullmatch(info.filename)
            if match:  # slides first, then notes, each in slide order
                parts.append((match[1].startswith("notes"), int(match[2]), info))
        if sum(info.file_size for *_, info in parts) > MAX_UNZIPPED_BYTES:
            return ""  # likely a zip bomb: report it as unreadable rather than unpack it
        lines, size = [], 0
        for *_, info in sorted(parts, key=lambda p: p[:2]):
            xml = archive.read(info)  # zipfile stops at the declared size and checks the CRC
            if b"<!DOCTYPE" in xml:
                continue  # Office never writes a DTD; one here is an entity-expansion trick
            for line in _drawing_lines(ElementTree.fromstring(xml)):
                lines.append(line)
                size += len(line)
                if size > MAX_CHARS:
                    return "\n".join(lines)
    return "\n".join(lines)


def _drawing_lines(element):
    """One line per paragraph, and one comma-separated line per table row."""
    if element.tag == _DRAWING + "tr":
        cells = (" ".join(_paragraph_text(p) for p in cell.iter(_DRAWING + "p"))
                 for cell in element.iter(_DRAWING + "tc"))
        yield ",".join(cells)
    elif element.tag == _DRAWING + "p":
        yield _paragraph_text(element)
    else:
        for child in element:
            yield from _drawing_lines(child)


def _paragraph_text(paragraph):
    return "".join(run.text or "" for run in paragraph.iter(_DRAWING + "t"))
